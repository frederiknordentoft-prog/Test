#!/usr/bin/env python3
"""verify_assets.py — GIF and PPTX checks for tests/verify.mjs.

Called by verify.mjs through child_process, but usable on its own:

    python3 tests/verify_assets.py gif  tests/out/anim.gif  --frames 15 --width 600 --height 384 --loop forever
    python3 tests/verify_assets.py pptx tests/out/native.pptx --png tests/out/pptx-slide1.png --slides 2 --min-shapes 8

Output protocol (one line per check, parsed by verify.mjs):

    PASS <name> — <detail>
    FAIL <name> — <reason>
    SKIP <name> — <reason>

Exit code is 1 when any check FAILed.

Requirements: Pillow, python-pptx. LibreOffice (soffice) is optional: when it
can render the PPTX we use it; otherwise slide 1 is drawn with Pillow straight
from the shape tree (blockArc / triangle / ellipse / roundRect / text), which
is enough to eyeball the geometry next to the SVG render.
"""

import argparse
import math
import pathlib
import shutil
import subprocess
import sys

EMU_PER_IN = 914400
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

_failed = False


def report(status, name, detail=""):
    """Print one protocol line and remember failures for the exit code."""
    global _failed
    if status == "FAIL":
        _failed = True
    line = f"{status} {name}"
    if detail:
        line += f" — {detail}"
    print(line, flush=True)


def check(cond, name, ok_detail="", bad_detail=None):
    report("PASS" if cond else "FAIL", name, ok_detail if cond else (bad_detail or ok_detail))
    return bool(cond)


# ---------------------------------------------------------------------------
# GIF
# ---------------------------------------------------------------------------

def verify_gif(args):
    from PIL import Image, ImageChops

    path = pathlib.Path(args.path)
    if not path.is_file():
        report("FAIL", "gif-exists", f"{path} not found")
        return
    try:
        im = Image.open(path)
    except Exception as exc:  # noqa: BLE001 — any decode failure is a test failure
        report("FAIL", "gif-decodes", f"Pillow could not open {path.name}: {exc}")
        return

    n_frames = getattr(im, "n_frames", 1)
    check(im.format == "GIF", "gif-decodes", f"{path.name}: format {im.format}, {n_frames} frames, {path.stat().st_size:,} bytes")

    # Every frame must decode; collect durations and a few frames for the motion check.
    durations, frames = [], {}
    wanted = {0, n_frames // 2, n_frames - 1}
    try:
        for i in range(n_frames):
            im.seek(i)
            durations.append(int(im.info.get("duration", 0)))
            if i in wanted:
                frames[i] = im.convert("RGB").copy()
        report("PASS", "gif-frames-decode", f"all {n_frames} frames decoded")
    except Exception as exc:  # noqa: BLE001
        report("FAIL", "gif-frames-decode", f"frame {len(durations)} failed: {exc}")
        return

    if args.frames is not None:
        check(n_frames == args.frames, "gif-frame-count", f"{n_frames} frames", f"expected {args.frames} frames, got {n_frames}")
    else:
        check(n_frames >= 2, "gif-frame-count", f"{n_frames} frames", "a GIF animation needs at least 2 frames")

    if args.width is not None and args.height is not None:
        check(im.size == (args.width, args.height), "gif-size", f"{im.size[0]}×{im.size[1]} px",
              f"expected {args.width}×{args.height}, got {im.size[0]}×{im.size[1]}")
    else:
        report("PASS", "gif-size", f"{im.size[0]}×{im.size[1]} px")

    # Loop metadata: Pillow exposes the NETSCAPE2.0 extension as info['loop']
    # (0 = forever). When the block is absent the GIF plays once.
    im.seek(0)
    loop = im.info.get("loop", None)
    if args.loop == "forever":
        check(loop == 0, "gif-loop", "loops forever (NETSCAPE loop=0)", f"expected loop=0, got {loop!r}")
    elif args.loop == "once":
        check(loop is None or loop == 1, "gif-loop", "plays once (no NETSCAPE loop block)" if loop is None else f"plays once (loop={loop})",
              f"expected no loop block / loop=1, got {loop!r}")
    else:
        report("PASS", "gif-loop", f"loop metadata: {loop!r}")

    # Timing: browsers clamp delays below ~20 ms, so anything in 20–200 ms is sane.
    if durations:
        avg = sum(durations) / len(durations)
        fps = 1000 / avg if avg else 0
        check(all(20 <= d <= 200 for d in durations), "gif-timing", f"{avg:.0f} ms/frame ≈ {fps:.0f} fps",
              f"frame delays outside 20–200 ms: {sorted(set(durations))}")

    # Motion: the needle must actually move — first, middle and last frame differ.
    if len(frames) >= 2 and n_frames >= 2:
        def diff_ratio(a, b):
            bbox = ImageChops.difference(a, b).getbbox()
            if not bbox:
                return 0.0
            area = (bbox[2] - bbox[0]) * (bbox[3] - bbox[1])
            return area / float(a.size[0] * a.size[1])

        first, last = frames[0], frames[n_frames - 1]
        r_fl = diff_ratio(first, last)
        mid = frames.get(n_frames // 2)
        r_fm = diff_ratio(first, mid) if mid is not None else r_fl
        check(r_fl > 0.001 and r_fm > 0.001, "gif-motion",
              f"frames differ (first↔last bbox {r_fl:.1%}, first↔mid {r_fm:.1%} of the image)",
              f"frames look identical (first↔last {r_fl:.2%}, first↔mid {r_fm:.2%}) — is the needle animating?")


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------

def _prst(shape):
    sppr = getattr(shape._element, "spPr", None)
    if sppr is None:
        return None
    prst = sppr.find(A_NS + "prstGeom")
    return prst.get("prst") if prst is not None else None


def _adj(shape, name, default=None):
    sppr = shape._element.spPr
    prst = sppr.find(A_NS + "prstGeom")
    if prst is None:
        return default
    for g in prst.iter(A_NS + "gd"):
        if g.get("name") == name:
            try:
                return int(g.get("fmla").split()[1])
            except (IndexError, ValueError):
                return default
    return default


def _solid_fill(node):
    """(rgb tuple, alpha 0–1) for an <a:solidFill> child of node, or (None, 0)."""
    if node is None:
        return None, 0.0
    sf = node.find(A_NS + "solidFill")
    if sf is None or len(sf) == 0:
        return None, 0.0
    clr = sf[0]
    val = clr.get("val")
    if clr.tag != A_NS + "srgbClr" or not val:
        return (128, 128, 128), 1.0  # scheme colour — approximate
    a = clr.find(A_NS + "alpha")
    alpha = int(a.get("val")) / 100000 if a is not None else 1.0
    return tuple(int(val[i:i + 2], 16) for i in (0, 2, 4)), alpha


def _line(shape):
    ln = shape._element.spPr.find(A_NS + "ln")
    if ln is None or ln.find(A_NS + "noFill") is not None:
        return None
    rgb, alpha = _solid_fill(ln)
    if rgb is None:
        return None
    dash = ln.find(A_NS + "prstDash")
    return {"rgb": rgb, "alpha": alpha, "w_pt": int(ln.get("w", 12700)) / 12700,
            "dash": dash.get("val") if dash is not None else "solid"}


def _rot_cw(x, y, deg):
    r = math.radians(deg)
    c, s = math.cos(r), math.sin(r)
    return (x * c - y * s, x * s + y * c)


def render_with_soffice(pptx_path, out_png, timeout=180):
    """Render slide 1 with LibreOffice. Returns (ok, detail)."""
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return False, "soffice not found on PATH"
    work = out_png.parent / ".soffice-tmp"
    shutil.rmtree(work, ignore_errors=True)
    work.mkdir(parents=True, exist_ok=True)
    profile = (work / "profile").resolve().as_uri()
    filters = [
        'png:impress_png_Export:{"PixelWidth":{"type":"long","value":"2000"},"PixelHeight":{"type":"long","value":"1125"}}',
        "png",
    ]
    last = ""
    try:
        for flt in filters:
            cmd = [exe, "--headless", "--norestore", f"-env:UserInstallation={profile}",
                   "--convert-to", flt, "--outdir", str(work), str(pptx_path)]
            try:
                proc = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
            except subprocess.TimeoutExpired:
                return False, f"soffice timed out after {timeout}s"
            produced = work / (pptx_path.stem + ".png")
            if produced.exists() and produced.stat().st_size > 0:
                shutil.copyfile(produced, out_png)
                return True, f"{pathlib.Path(exe).name} → {out_png.name}"
            msg = [l for l in (proc.stderr + "\n" + proc.stdout).splitlines() if l.strip() and "javaldx" not in l]
            last = msg[-1].strip() if msg else f"exit code {proc.returncode}"
        return False, f"soffice produced no PNG ({last})"
    finally:
        shutil.rmtree(work, ignore_errors=True)


def render_with_pillow(prs, out_png, dpi=200):
    """Approximate render of slide 1 from the shape tree (no LibreOffice needed)."""
    from PIL import Image, ImageDraw, ImageFont

    slide = prs.slides[0]
    sw, sh = prs.slide_width / EMU_PER_IN, prs.slide_height / EMU_PER_IN
    ss = 2  # supersample for smoother edges, then downscale
    S = dpi * ss
    W, H = int(round(sw * S)), int(round(sh * S))
    img = Image.new("RGBA", (W, H), (255, 255, 255, 255))

    def font(px):
        for f in ("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
                  "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
                  "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
                  "C:/Windows/Fonts/arialbd.ttf"):
            try:
                return ImageFont.truetype(f, int(px))
            except OSError:
                continue
        try:
            return ImageFont.load_default(size=int(px))
        except TypeError:  # Pillow < 10.1
            return ImageFont.load_default()

    def P(x, y):
        return (x * S, y * S)

    for shp in slide.shapes:
        if shp.width is None or shp.height is None:
            continue
        x, y = shp.left / EMU_PER_IN, shp.top / EMU_PER_IN
        w, h = shp.width / EMU_PER_IN, shp.height / EMU_PER_IN
        cx, cy = x + w / 2, y + h / 2
        prst = _prst(shp)
        sppr = getattr(shp._element, "spPr", None)
        rgb, alpha = _solid_fill(sppr)
        fill = (rgb + (int(round(alpha * 255)),)) if rgb and alpha > 0 else None
        ln = _line(shp)
        outline = (ln["rgb"] + (int(round(ln["alpha"] * 255)),)) if ln else None
        layer = Image.new("RGBA", (W, H), (0, 0, 0, 0))
        d = ImageDraw.Draw(layer)

        if prst == "blockArc":
            # OOXML: adj1/adj2 = start/end angle in 60000ths of a degree, clockwise from 3 o'clock
            a1 = _adj(shp, "adj1", 10800000) / 60000
            a2 = _adj(shp, "adj2", 0) / 60000
            a3 = _adj(shp, "adj3", 25000)
            R = min(w, h) / 2
            r_in = max(0.0, R - min(w, h) * a3 / 100000)
            sweep = (a2 - a1) % 360 or 360
            n = max(12, int(sweep / 2))
            pts = [P(cx + R * math.cos(math.radians(a1 + sweep * i / n)), cy + R * math.sin(math.radians(a1 + sweep * i / n))) for i in range(n + 1)]
            pts += [P(cx + r_in * math.cos(math.radians(a1 + sweep * i / n)), cy + r_in * math.sin(math.radians(a1 + sweep * i / n))) for i in range(n, -1, -1)]
            if fill:
                d.polygon(pts, fill=fill)
        elif prst == "triangle":
            # apex top-centre, base at the bottom; rotation is clockwise about the centre
            rot = shp.rotation or 0.0
            local = [(0, -h / 2), (w / 2, h / 2), (-w / 2, h / 2)]
            pts = [P(cx + _rot_cw(px, py, rot)[0], cy + _rot_cw(px, py, rot)[1]) for px, py in local]
            if fill or outline:
                d.polygon(pts, fill=fill, outline=outline, width=max(1, int(ln["w_pt"] / 72 * S)) if ln else 0)
        elif prst == "ellipse":
            d.ellipse([P(x, y), P(x + w, y + h)], fill=fill, outline=outline, width=max(1, int(ln["w_pt"] / 72 * S)) if ln else 0)
        elif prst in ("roundRect", "rect") or shp.has_text_frame:
            if fill or outline:
                radius = (h * S / 2) if prst == "roundRect" else 0
                d.rounded_rectangle([P(x, y), P(x + w, y + h)], radius=radius, fill=fill, outline=outline,
                                    width=max(1, int(ln["w_pt"] / 72 * S)) if ln else 0)
        elif shp.shape_type == 13:  # picture
            d.rectangle([P(x, y), P(x + w, y + h)], outline=(180, 180, 180, 255), width=2)

        if shp.has_text_frame and shp.text_frame.text.strip():
            para = shp.text_frame.paragraphs[0]
            run = para.runs[0] if para.runs else None
            size_pt = (run.font.size.pt if run is not None and run.font.size else 18)
            colour = (0, 0, 0)
            try:
                if run is not None and run.font.color and run.font.color.type is not None and run.font.color.rgb is not None:
                    colour = tuple(run.font.color.rgb)
            except (AttributeError, TypeError):
                pass
            d.text(P(cx, cy), shp.text_frame.text.strip(), fill=colour + (255,), font=font(size_pt / 72 * S), anchor="mm")

        img = Image.alpha_composite(img, layer)

    img = img.resize((W // ss, H // ss), Image.LANCZOS).convert("RGB")
    img.save(out_png)
    return img.size


def verify_pptx(args):
    from pptx import Presentation

    path = pathlib.Path(args.path)
    if not path.is_file():
        report("FAIL", "pptx-exists", f"{path} not found")
        return
    try:
        prs = Presentation(str(path))
    except Exception as exc:  # noqa: BLE001
        report("FAIL", "pptx-opens", f"python-pptx could not open {path.name}: {exc}")
        return

    sw, sh = prs.slide_width / EMU_PER_IN, prs.slide_height / EMU_PER_IN
    check(abs(sw - 10) < 0.01 and abs(sh - 5.625) < 0.01, "pptx-opens",
          f"{path.name}: {len(prs.slides)} slides, {sw:.3f} × {sh:.3f} in (16:9), {path.stat().st_size:,} bytes",
          f"slide size {sw:.3f} × {sh:.3f} in — expected 10 × 5.625 in (LAYOUT_16x9)")

    check(len(prs.slides) == args.slides, "pptx-slide-count", f"{len(prs.slides)} slides",
          f"expected {args.slides} slides, got {len(prs.slides)}")
    if len(prs.slides) == 0:
        return

    # Slide 1: native shapes
    slide = prs.slides[0]
    counts = {}
    names = []
    for shp in slide.shapes:
        kind = _prst(shp) or ("picture" if shp.shape_type == 13 else str(shp.shape_type))
        counts[kind] = counts.get(kind, 0) + 1
        names.append(shp.name)
    total = len(slide.shapes)
    breakdown = ", ".join(f"{k}×{v}" for k, v in sorted(counts.items(), key=lambda kv: -kv[1]))
    check(total >= args.min_shapes, "pptx-slide1-shapes", f"{total} shapes ({breakdown})",
          f"only {total} shapes on slide 1 ({breakdown}); expected at least {args.min_shapes}")
    check(counts.get("blockArc", 0) >= 3, "pptx-slide1-arcs", f"{counts.get('blockArc', 0)} blockArc segments (3 zones + sweep wedge when the trail is on)",
          f"expected ≥ 3 blockArc shapes for red/amber/green, found {counts.get('blockArc', 0)}")
    check(counts.get("triangle", 0) >= 1, "pptx-slide1-needle", f"{counts.get('triangle', 0)} triangle shapes (needle, ghosts, prev marker)",
          "no triangle shape — the needle is missing")
    check(counts.get("ellipse", 0) >= 2, "pptx-slide1-hub", f"{counts.get('ellipse', 0)} ellipses (hub disc / ring / core)",
          f"expected ≥ 2 ellipses for the hub, found {counts.get('ellipse', 0)}")
    named = sum(1 for n in names if n and not n.lower().startswith(("shape", "autoshape", "object")))
    check(named >= max(1, total // 2), "pptx-slide1-names", f"{named}/{total} shapes carry descriptive names (objectName)",
          f"only {named}/{total} shapes are named — hard to edit in PowerPoint's selection pane")

    # Slide 2: the PNG image
    if len(prs.slides) >= 2:
        pics = [s for s in prs.slides[1].shapes if s.shape_type == 13]
        check(len(pics) >= 1, "pptx-slide2-image", f"{len(pics)} picture on slide 2" + (
            f" ({pics[0].width / EMU_PER_IN:.2f} × {pics[0].height / EMU_PER_IN:.2f} in)" if pics else ""),
              "slide 2 has no picture — the PNG fallback slide is missing")

    # Render slide 1 → PNG (LibreOffice if it works here, otherwise Pillow)
    if args.png:
        out_png = pathlib.Path(args.png)
        out_png.parent.mkdir(parents=True, exist_ok=True)
        ok, detail = render_with_soffice(path, out_png)
        if ok:
            from PIL import Image
            with Image.open(out_png) as im:
                report("PASS", "pptx-render", f"{detail} ({im.size[0]}×{im.size[1]} px, LibreOffice)")
        else:
            report("SKIP", "pptx-render-soffice", f"LibreOffice could not render it here ({detail}); drawing slide 1 with Pillow instead")
            try:
                size = render_with_pillow(prs, out_png)
                report("PASS", "pptx-render-fallback", f"{out_png.name} ({size[0]}×{size[1]} px) drawn from the shape tree — compare with svg-reference.png")
            except Exception as exc:  # noqa: BLE001
                report("FAIL", "pptx-render-fallback", f"Pillow render failed: {exc}")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    g = sub.add_parser("gif", help="check an animated GIF")
    g.add_argument("path")
    g.add_argument("--frames", type=int, default=None, help="expected frame count")
    g.add_argument("--width", type=int, default=None)
    g.add_argument("--height", type=int, default=None)
    g.add_argument("--loop", choices=["forever", "once", "any"], default="any")
    g.set_defaults(func=verify_gif)

    p = sub.add_parser("pptx", help="check a PPTX and render slide 1")
    p.add_argument("path")
    p.add_argument("--png", default=None, help="where to write the slide-1 render")
    p.add_argument("--slides", type=int, default=2, help="expected slide count")
    p.add_argument("--min-shapes", type=int, default=8, help="minimum shapes on slide 1")
    p.set_defaults(func=verify_pptx)

    args = ap.parse_args(argv)
    try:
        args.func(args)
    except ImportError as exc:
        report("FAIL", f"{args.cmd}-deps", f"missing Python dependency: {exc} (pip install Pillow python-pptx)")
    return 1 if _failed else 0


if __name__ == "__main__":
    sys.exit(main())
